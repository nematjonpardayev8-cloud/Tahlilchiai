from __future__ import annotations

from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill


def export_excel(df: pd.DataFrame) -> bytes:
    stream = BytesIO()
    df.to_excel(stream, index=False, sheet_name="Audit Results")
    stream.seek(0)
    wb = load_workbook(stream)
    ws = wb["Audit Results"]
    headers = {cell.value: cell.column for cell in ws[1]}
    fill = PatternFill("solid", fgColor="FFC7CE")
    for row in range(2, ws.max_row + 1):
        score = ws.cell(row, headers.get("risk_score", 1)).value or 0
        if score >= 70:
            for cell in ws[row]:
                cell.fill = fill
            ws.cell(row, 1).comment = Comment("High-risk transaction requires audit investigation.", "Audit Sentinel")
    result = BytesIO(); wb.save(result); return result.getvalue()
